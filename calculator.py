"""Read saved Excel calculation values, never pivot tables or execute formulas."""
from io import BytesIO
from math import isfinite
from decimal import Decimal
import re
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

GROUPS = ('FREZ', 'LASER', 'MANUALNE', 'SZTANCA', 'GILOTYNA', 'FORMIERKA')
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def read_materials(ws, root, quantity):
    """Pair each material header with its own total column; retain separate blocks."""
    cells = {cell.coordinate: cell for row in ws.iter_rows() for cell in row
             if cell.value is not None}
    def normalized(value):
        return ' '.join(str(value or '').casefold().replace('²', '2').split())

    headers = [c for c in cells.values() if normalized(c.value) == 'tworzywo i opis']
    totals = [c for c in cells.values() if normalized(c.value) == 'm2 łącznie']
    merges = [openpyxl.worksheet.cell_range.CellRange(m.attrib['ref'])
              for m in root.findall('m:mergeCells/m:mergeCell', NS)]
    def bottom(cell):
        return next((m.max_row for m in merges if cell.coordinate in m), cell.row)

    materials, warnings = [], []
    for header in headers:
        right = min((h.column for h in headers if h.row == header.row and h.column > header.column),
                    default=ws.max_column + 1)
        end = min((h.row for h in headers if h.column == header.column and h.row > header.row),
                  default=ws.max_row + 1)
        # Stop before the process section, even for the last material block.
        end = min([end] + [c.row for c in cells.values()
                          if c.row > header.row and normalized(c.value) == 'opis procesu'])
        description_row = bottom(header) + 1
        description = cells.get(f'{header.column_letter}{description_row}')
        if description is None or not str(description.value).strip() or description_row >= end:
            continue
        label = str(description.value).strip()
        matches = [c for c in totals if header.row <= c.row <= description_row
                   and header.column < c.column < right]
        if len(matches) != 1:
            warnings.append(f'{header.coordinate}: {label} — nie znaleziono jednoznacznego nagłówka m2 łącznie.')
            continue
        total_header = matches[0]
        address = f'{total_header.column_letter}{bottom(total_header) + 1}'
        value = cells.get(address)
        value = value.value if value is not None else None
        if (isinstance(value, bool) or not isinstance(value, (int, float))
                or not isfinite(value) or value < 0):
            warnings.append(f'{address}: {label} — brak poprawnego m² łącznie. Przelicz i zapisz plik w Excelu.')
            continue
        total = Decimal(str(value))
        materials.append({'description': label, 'total_m2': float(total),
                          'per_item': float(total / Decimal(str(quantity))),
                          'description_cell': description.coordinate, 'total_cell': address})
    return materials, warnings


def read_workbook(data, sheet=None):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError('Rozpakowany plik jest zbyt duży (limit 50 MB).')
    book = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    try:
        candidates = []
        for ws in book:
            if ws.title.upper() == 'TAB':
                continue
            for row in ws.iter_rows(max_row=min(ws.max_row, 100), max_col=13, values_only=True):
                if str(row[7]).strip().upper() == 'OPIS PROCESU' and 'RBH' in str(row[12]).upper():
                    candidates.append(ws.title)
                    break
        if not candidates:
            raise ValueError('Nie znaleziono arkusza wyceny z nagłówkami OPIS PROCESU (H) i ilość RBH (M).')
        if sheet is None and len(candidates) > 1:
            return {'sheets': candidates, 'selection_required': True}
        selected = sheet or candidates[0]
        if selected not in candidates:
            raise ValueError('Wybrany arkusz nie jest arkuszem wyceny.')
        ws = book[selected]
        quantity = ws['K4'].value
        if isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or not isfinite(quantity) or quantity <= 0 or quantity != int(quantity):
            raise ValueError('Nakład w K4 musi być dodatnią całkowitą liczbą sztuk. Zapisz przeliczony plik w Excelu.')
        # Read raw numeric day fractions: datetime conversion rounds subsecond values.
        with zipfile.ZipFile(BytesIO(data)) as archive:
            root = ET.fromstring(archive.read(ws._worksheet_path))
        raw = {c.attrib['r']: c.find('m:v', NS).text for c in root.findall('.//m:c', NS)
               if c.find('m:v', NS) is not None and c.attrib.get('t', 'n') == 'n'}
        totals = dict.fromkeys(GROUPS, Decimal(0))
        details = []
        started = False
        setup = Decimal(0)
        for row in ws.iter_rows(max_col=13):
            label = str(row[7].value or '').strip()
            if label.upper() == 'OPIS PROCESU':
                started = True
                continue
            if not started:
                continue
            if not label:
                break
            group = re.split(r'\s*[-–—]\s*', label, maxsplit=1)[0].strip().upper()
            is_setup = group == 'RBH NA ROZRUCH PRODUKCJI'
            if is_setup:
                group = 'MANUALNE'
            if group not in totals:
                raise ValueError(f'Nieznana grupa procesu w H{row[7].row}: {label}. Wynik nie został obliczony.')
            address = row[12].coordinate
            try:
                hours = Decimal(raw[address]) * 24
                if not hours.is_finite() or hours < 0:
                    raise ValueError()
            except (KeyError, ValueError, ArithmeticError, TypeError):
                raise ValueError(f'Brak poprawnego czasu w {address}. Otwórz plik w Excelu, przelicz formuły i zapisz ponownie.') from None
            totals[group] += hours
            if is_setup:
                setup += hours
            details.append({'row': row[7].row, 'label': label, 'group': group, 'hours': float(hours)})
        if not details:
            raise ValueError('Tabela procesów jest pusta.')
        materials, material_warnings = read_materials(ws, root, quantity)
        return {'sheets': candidates, 'sheet': selected, 'quantity': int(quantity),
                'project': str(ws['H4'].value or ''), 'setup_hours': float(setup),
                'processes': [{'name': g, 'hours': float(h), 'per_item': float(h / Decimal(str(quantity)))} for g, h in totals.items()],
                'details': details, 'materials': materials, 'material_warnings': material_warnings}
    finally:
        book.close()
