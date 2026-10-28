import unittest
from io import BytesIO
import openpyxl
from calculator import read_workbook, summarize_materials


def workbook(quantity=10, label='FREZ - cięcie', time=1.5, extra=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Wycena'
    sheet['K4'] = quantity
    sheet['H9'] = 'OPIS PROCESU'
    sheet['M9'] = 'ilość RBH'
    sheet['H10'] = label
    sheet['M10'] = time
    sheet['H11'] = 'RBH NA ROZRUCH PRODUKCJI'
    sheet['M11'] = 2 / 24
    if extra:
        book.copy_worksheet(sheet).title = 'Druga'
    out = BytesIO()
    book.save(out)
    return out.getvalue()


class CalculatorTests(unittest.TestCase):
    def additions_workbook(self, divisor=12, amount=1, foam=.12):
        book = openpyxl.load_workbook(BytesIO(workbook(quantity=8000)))
        ws = book.active
        for r in (43, 60):
            ws.cell(r, 3, 'ILOŚĆ SZT. KOMPONENTNU')
            ws.cell(r, 5, 'SPECYFIKACJA')
            ws.cell(r, 8, 'OPIS')
        ws['K43'] = 'DODATEK NA ILOŚĆ SZTUK Z NAKŁADU'
        for r, c, value in [(44, 3, amount), (44, 5, '230x280x440'),
                            (44, 8, 'karton na 12'), (44, 11, divisor),
                            (46, 3, 6), (46, 5, 'Nit + podkładka'), (46, 8, 'Wojtech'), (46, 11, 1),
                            (54, 5, 'm2 pianki'), (54, 6, 'OPIS'),
                            (55, 5, foam), (55, 6, 'Pianka ochronna'), (55, 11, 2),
                            (58, 2, 'DODATKI NA CAŁY NAKŁAD + DOSTAWY TWORZYWA'),
                            (61, 3, 2), (61, 5, 'wykrojnik kieszonki x2'),
                            (67, 2, 'DAP - TRANSPORT DO KLIENTA'),
                            (70, 3, 5), (70, 5, 'Nie jest dodatkiem')]:
            ws.cell(r, c, value)
        # Deliberately unrelated costs must never become consumption quantities.
        ws['J44'] = 2000
        ws['L44'] = .25
        out = BytesIO()
        book.save(out)
        book.close()
        return out.getvalue()

    def test_additions_quantities_and_section_boundaries(self):
        result = read_workbook(self.additions_workbook())
        items = result['additions']
        self.assertEqual([a['row'] for a in items], [44, 46, 55, 61])
        self.assertAlmostEqual(items[0]['per_item'], 1 / 12)
        self.assertEqual(items[1]['per_item'], 6)
        self.assertEqual(items[2]['per_item'], .06)
        self.assertEqual(items[2]['unit'], 'm²')
        self.assertEqual(items[3]['per_item'], .00025)
        self.assertEqual(result['addition_warnings'], [])
        self.assertEqual(result['processes'], read_workbook(workbook(quantity=8000))['processes'])

    def test_invalid_component_divisor_does_not_hide_other_results(self):
        for divisor in (None, 0, -1, True, '12'):
            with self.subTest(divisor=divisor):
                result = read_workbook(self.additions_workbook(divisor=divisor))
                self.assertEqual(len(result['additions']), 3)
                self.assertIn('K44', result['addition_warnings'][0])
                self.assertTrue(result['processes'])

    def test_invalid_addition_quantity_or_formula_cache(self):
        for amount in (None, -1, True, '=1+1', '#VALUE!'):
            with self.subTest(amount=amount):
                result = read_workbook(self.additions_workbook(amount=amount))
                self.assertIn('C44', result['addition_warnings'][0])
        result = read_workbook(self.additions_workbook(foam='=C55*D55/10000'))
        self.assertIn('E55', result['addition_warnings'][0])

    def test_zero_quantity_and_absent_additions(self):
        self.assertEqual(read_workbook(self.additions_workbook(amount=0))['additions'][0]['per_item'], 0)
        self.assertEqual(read_workbook(workbook())['additions'], [])
        self.assertEqual(read_workbook(workbook())['addition_warnings'], [])

    def test_repeated_materials_are_summed_before_rounding(self):
        result = read_workbook(self.material_workbook(
            labels=('OPAL SATYNA 3mm WH 10 DC - płytka', 'OPAL SATYNA 3mm WH 10 DC - litery'),
            values=(.809053, .437675)))
        self.assertEqual(len(result['materials']), 2)
        self.assertEqual(len(result['material_totals']), 1)
        group = result['material_totals'][0]
        self.assertEqual(group['description'], 'OPAL SATYNA 3mm WH 10 DC')
        self.assertEqual(group['count'], 2)
        self.assertEqual(group['total_m2'], 1.246728)
        self.assertEqual(group['per_item'], .1246728)

    def test_equivalent_material_spelling_and_word_order(self):
        for kind in ('PMMA', 'HIPS', 'PC', 'PET', 'Błona'):
            with self.subTest(kind=kind):
                rows = [{'description': f'{kind} 3mm biała - płytka', 'total_m2': .1},
                        {'description': f'  {kind.lower()} WH 3,00 mm — litery', 'total_m2': .2}]
                groups = summarize_materials(rows, 10)
                self.assertEqual(len(groups), 1)
                self.assertEqual(groups[0]['total_m2'], .3)
                self.assertEqual(groups[0]['per_item'], .03)

    def test_distinct_material_type_thickness_color_finish_model(self):
        names = ['PMMA 3mm WH - płytka', 'HIPS 3mm WH - litery',
                 'PMMA 4mm WH - litery', 'PMMA 3mm BK - litery',
                 'PMMA 3mm WH satyna - litery', 'PMMA 3mm WH 10 DC - litery',
                 'Błona 3M 7955MP - płytka', 'Błona 3M 467MP - litery']
        self.assertEqual(len(summarize_materials(
            [{'description': n, 'total_m2': 1} for n in names], 10)), len(names))

    def test_identical_names_without_dash_and_empty_materials(self):
        groups = summarize_materials([{'description': 'PET 1mm clear', 'total_m2': 0},
                                     {'description': 'PET 1 mm bezbarwny', 'total_m2': 2}], 4)
        self.assertEqual(groups[0]['count'], 2)
        self.assertEqual(groups[0]['per_item'], .5)
        self.assertEqual(summarize_materials([], 10), [])

    def test_production_layout_h75_h76_j76_j77(self):
        book = openpyxl.load_workbook(BytesIO(workbook(quantity=250)))
        ws = book.active
        for offset, label, total in ((0, 'PVC biały 3 mm', 62.5), (10, 'PET', 20),
                                     (20, '   ', 99)):
            ws.cell(75 + offset, 8, 'TWORZYWO I OPIS')
            ws.cell(76 + offset, 8, label)
            ws.cell(76 + offset, 10, 'm2 łącznie')
            ws.cell(77 + offset, 10, total)
        out = BytesIO()
        book.save(out)
        book.close()
        result = read_workbook(out.getvalue())
        self.assertEqual([m['per_item'] for m in result['materials']], [.25, .08])
        self.assertEqual(result['materials'][0]['description_cell'], 'H76')
        self.assertEqual(result['materials'][0]['total_cell'], 'J77')
        self.assertEqual(result['material_warnings'], [])
        self.assertEqual(result['processes'], read_workbook(workbook(quantity=250))['processes'])

    def material_workbook(self, values=(25, 0), labels=('PVC 3 mm', 'PET')):
        book = openpyxl.load_workbook(BytesIO(workbook()))
        ws = book.active
        ws.move_range('H9:M11', rows=40)
        for row, total, label in zip((9, 20), values, labels):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.cell(row, 1, 'TWORZYWO I OPIS')
            ws.cell(row + 1, 1, label)
            ws.cell(row, 6, 'm²\nłącznie')
            ws.cell(row + 1, 6, total)
        out = BytesIO()
        book.save(out)
        book.close()
        return out.getvalue()

    def test_material_blocks_and_process_regression(self):
        result = read_workbook(self.material_workbook())
        self.assertEqual([m['description'] for m in result['materials']], ['PVC 3 mm', 'PET'])
        self.assertEqual([m['per_item'] for m in result['materials']], [2.5, 0])
        self.assertEqual(result['material_warnings'], [])
        original = read_workbook(workbook())
        self.assertEqual(result['processes'], original['processes'])
        self.assertEqual(result['setup_hours'], original['setup_hours'])

    def test_empty_description_skips_block(self):
        result = read_workbook(self.material_workbook(labels=('   ', 'PET')))
        self.assertEqual(len(result['materials']), 1)
        self.assertEqual(result['materials'][0]['description'], 'PET')

    def test_missing_or_invalid_material_total_preserves_processes(self):
        for value in (None, '=25/2', '#DIV/0!', -1, True, '25'):
            with self.subTest(value=value):
                result = read_workbook(self.material_workbook(values=(value, 12.5)))
                self.assertEqual(len(result['materials']), 1)
                self.assertEqual(result['materials'][0]['per_item'], 1.25)
                self.assertEqual(len(result['material_warnings']), 1)
                self.assertIn('F10', result['material_warnings'][0])
                self.assertEqual(result['processes'], read_workbook(workbook())['processes'])

    def test_old_workbook_has_no_materials(self):
        result = read_workbook(workbook())
        self.assertEqual(result['materials'], [])
        self.assertEqual(result['material_warnings'], [])

    def test_adjacent_blocks_and_duplicate_descriptions(self):
        book = openpyxl.load_workbook(BytesIO(self.material_workbook()))
        ws = book.active
        ws['H20'] = 'TWORZYWO I OPIS'
        ws['H21'] = 'PET'
        ws['M20'] = 'm2 łącznie'
        ws['M21'] = 15
        out = BytesIO()
        book.save(out)
        book.close()
        result = read_workbook(out.getvalue())
        self.assertEqual([m['per_item'] for m in result['materials']], [2.5, 0, 1.5])

    def test_vertical_merged_headers(self):
        book = openpyxl.load_workbook(BytesIO(self.material_workbook()))
        ws = book.active
        ws['A10'] = None
        ws['A11'] = 'PVC 3 mm'
        ws.unmerge_cells('A9:D9')
        ws.merge_cells('A9:D10')
        ws['F10'] = None
        ws['F11'] = 25
        ws.merge_cells('F9:F10')
        out = BytesIO()
        book.save(out)
        book.close()
        self.assertEqual(read_workbook(out.getvalue())['materials'][0]['per_item'], 2.5)

    def test_days_over_24_hours_and_setup(self):
        result = read_workbook(workbook())
        values = {p['name']: p for p in result['processes']}
        self.assertAlmostEqual(values['FREZ']['hours'], 36)
        self.assertAlmostEqual(values['FREZ']['per_item'], 3.6)
        self.assertAlmostEqual(values['MANUALNE']['per_item'], .2)
        self.assertEqual(len(result['details']), 2)

    def test_invalid_quantity(self):
        for quantity in (0, -1, 'dziesięć', 1.5, True):
            with self.subTest(quantity=quantity), self.assertRaises(ValueError):
                read_workbook(workbook(quantity=quantity))

    def test_unknown_process_blocks_results(self):
        with self.assertRaisesRegex(ValueError, 'Nieznana grupa'):
            read_workbook(workbook(label='INNY - opis'))

    def test_missing_formula_cache_blocks_results(self):
        with self.assertRaisesRegex(ValueError, 'M10'):
            read_workbook(workbook(time='=1/24'))

    def test_negative_time(self):
        with self.assertRaises(ValueError):
            read_workbook(workbook(time=-1))

    def test_multiple_sheets(self):
        data = workbook(extra=True)
        self.assertTrue(read_workbook(data)['selection_required'])
        self.assertEqual(read_workbook(data, 'Druga')['sheet'], 'Druga')

    def test_dash_variants(self):
        for dash in ('-', '–', '—'):
            self.assertEqual(read_workbook(workbook(label='FREZ '+dash+' opis'))['processes'][0]['per_item'], 3.6)


if __name__ == '__main__':
    unittest.main()
